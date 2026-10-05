import os
import re
import uuid
from pathlib import Path

from database import (
    CONOCIMIENTO_DIR,
    crear_fuente_conocimiento,
    guardar_fragmentos_conocimiento,
    marcar_fuente_procesada,
)


# ============================================================
# CONFIGURACIÓN
# ============================================================

EXTENSIONES_DOCUMENTOS = {".pdf", ".docx", ".txt"}
EXTENSIONES_IMAGENES = {".png", ".jpg", ".jpeg", ".webp"}

EXTENSIONES_PERMITIDAS = EXTENSIONES_DOCUMENTOS | EXTENSIONES_IMAGENES


# ============================================================
# UTILIDADES
# ============================================================

def limpiar_texto(texto):
    """
    Limpia espacios, tabulaciones, caracteres nulos y saltos
    de línea excesivos.
    """
    texto = (texto or "")

    texto = (
        texto
        .replace("\x00", " ")
        .replace("\r\n", "\n")
        .replace("\r", "\n")
    )

    # Espacios y tabulaciones repetidos
    texto = re.sub(r"[ \t]+", " ", texto)

    # Más de dos saltos de línea
    texto = re.sub(r"\n{3,}", "\n\n", texto)

    return texto.strip()


def es_imagen(extension):
    return extension.lower() in EXTENSIONES_IMAGENES


# ============================================================
# EXTRACCIÓN DE TEXTO
# ============================================================

def extraer_texto(ruta):
    """
    Extrae texto de PDF, DOCX y TXT.

    Las imágenes se procesan mediante la descripción introducida
    por el administrador y no mediante OCR automático.
    """

    ext = Path(ruta).suffix.lower()

    # --------------------------------------------------------
    # PDF
    # --------------------------------------------------------

    if ext == ".pdf":

        from pypdf import PdfReader

        lector = PdfReader(ruta)

        partes = []

        for numero, pagina in enumerate(lector.pages, 1):

            texto = limpiar_texto(
                pagina.extract_text() or ""
            )

            if texto:

                partes.append(
                    f"[Página {numero}]\n{texto}"
                )

        if not partes:

            raise ValueError(
                "El PDF no contiene texto extraíble. "
                "Si el documento está escaneado, deberá utilizar "
                "OCR antes de incorporarlo a la Base de conocimiento."
            )

        return "\n\n".join(partes)

    # --------------------------------------------------------
    # WORD
    # --------------------------------------------------------

    if ext == ".docx":

        from docx import Document

        documento = Document(ruta)

        partes = []

        # Párrafos
        for parrafo in documento.paragraphs:

            texto = limpiar_texto(parrafo.text)

            if texto:

                partes.append(texto)

        # Tablas
        for tabla in documento.tables:

            for fila in tabla.rows:

                valores = []

                for celda in fila.cells:

                    texto = limpiar_texto(celda.text)

                    if texto:

                        valores.append(texto)

                if valores:

                    partes.append(
                        " | ".join(valores)
                    )

        if not partes:

            raise ValueError(
                "El documento Word no contiene texto extraíble."
            )

        return "\n\n".join(partes)

    # --------------------------------------------------------
    # TXT
    # --------------------------------------------------------

    if ext == ".txt":

        codificaciones = (
            "utf-8",
            "utf-8-sig",
            "latin-1",
            "cp1252",
        )

        for codificacion in codificaciones:

            try:

                with open(
                    ruta,
                    "r",
                    encoding=codificacion
                ) as archivo:

                    texto = limpiar_texto(
                        archivo.read()
                    )

                    if not texto:

                        raise ValueError(
                            "El archivo TXT está vacío."
                        )

                    return texto

            except UnicodeDecodeError:

                continue

        raise ValueError(
            "No se pudo determinar la codificación del archivo TXT."
        )

    # --------------------------------------------------------
    # IMÁGENES
    # --------------------------------------------------------

    if ext in EXTENSIONES_IMAGENES:

        raise ValueError(
            "Las imágenes deben procesarse utilizando "
            "la descripción proporcionada por el administrador."
        )

    raise ValueError(
        "Formato de archivo no permitido."
    )


# ============================================================
# DIVIDIR TEXTO PARA RAG
# ============================================================

def dividir_texto(
    texto,
    tamano=1400,
    solapamiento=220
):

    texto = limpiar_texto(texto)

    if not texto:

        return []

    fragmentos = []

    inicio = 0

    longitud = len(texto)

    while inicio < longitud:

        fin = min(
            inicio + tamano,
            longitud
        )

        trozo = texto[inicio:fin]

        # Intentar cortar por párrafo o frase
        if fin < longitud:

            corte_parrafo = trozo.rfind("\n\n")
            corte_frase = trozo.rfind(". ")

            corte = max(
                corte_parrafo,
                corte_frase
            )

            if corte > tamano * 0.55:

                fin = inicio + corte + 1

                trozo = texto[inicio:fin]

        trozo = trozo.strip()

        if trozo:

            fragmentos.append(
                {
                    "contenido": trozo,
                    "metadata": (
                        f"caracteres:{inicio}-{fin}"
                    ),
                }
            )

        if fin >= longitud:

            break

        inicio = max(
            fin - solapamiento,
            inicio + 1
        )

    return fragmentos


# ============================================================
# PROCESAMIENTO PRINCIPAL
# ============================================================

def procesar_archivo_subido(
    archivo_subido,
    titulo,
    descripcion="",
    categoria="General",
    nivel_id=None,
    unidad_id=None,
    programa_id=None,
    creado_por=None,
    contenido_imagen="",
):

    ruta = None
    fuente_id = None

    try:

        # ----------------------------------------------------
        # VALIDAR ARCHIVO
        # ----------------------------------------------------

        if archivo_subido is None:

            raise ValueError(
                "Debes seleccionar un documento o imagen."
            )

        nombre_original = archivo_subido.name

        ext = Path(
            nombre_original
        ).suffix.lower()

        if ext not in EXTENSIONES_PERMITIDAS:

            raise ValueError(
                "Formato no permitido. "
                "Puedes utilizar PDF, DOCX, TXT, PNG, JPG, JPEG o WEBP."
            )

        # ----------------------------------------------------
        # CREAR DIRECTORIO
        # ----------------------------------------------------

        os.makedirs(
            CONOCIMIENTO_DIR,
            exist_ok=True
        )

        # ----------------------------------------------------
        # GUARDAR ARCHIVO
        # ----------------------------------------------------

        nombre_interno = (
            f"{uuid.uuid4().hex}{ext}"
        )

        ruta = os.path.join(
            CONOCIMIENTO_DIR,
            nombre_interno
        )

        with open(
            ruta,
            "wb"
        ) as archivo:

            archivo.write(
                archivo_subido.getbuffer()
            )

        # ----------------------------------------------------
        # PROCESAR CONTENIDO
        # ----------------------------------------------------

        if es_imagen(ext):

            # Para flyers e imágenes utilizamos la información
            # proporcionada por el administrador.

            contenido = limpiar_texto(
                contenido_imagen
            )

            descripcion_limpia = limpiar_texto(
                descripcion
            )

            titulo_limpio = limpiar_texto(
                titulo
            )

            if not contenido:

                # Permitimos utilizar la descripción como contenido
                contenido = descripcion_limpia

            if not contenido:

                raise ValueError(
                    "Para incorporar una imagen a UPRI debes escribir "
                    "el texto o la información relevante que contiene "
                    "el flyer."
                )

            texto = (
                f"Título: {titulo_limpio}\n"
                f"Categoría: {categoria}\n"
                f"Tipo de fuente: Imagen institucional\n\n"
                f"{contenido}"
            )

            tipo_archivo = "imagen"

        else:

            texto = extraer_texto(
                ruta
            )

            tipo_archivo = ext[1:]

        # ----------------------------------------------------
        # CREAR FRAGMENTOS
        # ----------------------------------------------------

        fragmentos = dividir_texto(
            texto
        )

        if not fragmentos:

            raise ValueError(
                "No se encontró contenido suficiente "
                "para incorporar a la Base de conocimiento."
            )

        # ----------------------------------------------------
        # REGISTRAR FUENTE
        # ----------------------------------------------------

        ok, mensaje, fuente_id = crear_fuente_conocimiento(
            titulo,
            descripcion,
            categoria,
            tipo_archivo,
            ruta,
            nombre_original,
            nivel_id,
            unidad_id,
            programa_id,
            creado_por,
        )

        if not ok:

            raise ValueError(
                mensaje
            )

        # ----------------------------------------------------
        # GUARDAR FRAGMENTOS
        # ----------------------------------------------------

        guardar_fragmentos_conocimiento(
            fuente_id,
            fragmentos
        )

        # ----------------------------------------------------
        # RESPUESTA
        # ----------------------------------------------------

        if es_imagen(ext):

            mensaje_final = (
                "Imagen incorporada correctamente "
                "a la Base de conocimiento de UPRI."
            )

        else:

            mensaje_final = (
                "Documento procesado correctamente."
            )

        return {
            "ok": True,
            "mensaje": mensaje_final,
            "fuente_id": fuente_id,
            "fragmentos": len(fragmentos),
            "tipo": tipo_archivo,
            "archivo": nombre_original,
        }

    # ========================================================
    # ERROR
    # ========================================================

    except Exception as e:

        if fuente_id:

            try:

                marcar_fuente_procesada(
                    fuente_id,
                    0
                )

            except Exception:

                pass

        elif ruta and os.path.exists(ruta):

            try:

                os.remove(ruta)

            except OSError:

                pass

        return {
            "ok": False,
            "mensaje": str(e),
            "fragmentos": 0,
        }